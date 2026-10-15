"""用例导出：对齐《功能测试提示词 v3.0》表格列。"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import OUTPUT_DIR
from ..models import TestCase


# 功能测试提示词 v3.0 — 列顺序固定
V3_HEADERS = [
    "用例编号",
    "系统名字",
    "模块分类",
    "一级标题",
    "二级标题",
    "设计需求",
    "测试追踪",
    "终止条件",
    "用例标题",
    "前提与约束",
    "输入及操作说明",
    "期望测试结果",
    "测试执行人",
    "测试时间",
    "实际测试结果",
    "未通过详情",
    "审核人",
    "审核时间",
    "备注",
]


def _cell(v: str) -> str:
    s = (v or "").replace("\r\n", " ").replace("\n", " ").replace("\r", " ").strip()
    return s if s else "-"


def _row(c: TestCase) -> list[str]:
    return [
        _cell(c.case_no or c.id),
        _cell(c.sys_name),
        _cell(c.module_cat or c.module),
        _cell(c.title_l1),
        _cell(c.title_l2),
        _cell(c.design_req),
        _cell(c.test_track),
        _cell(c.stop_cond),
        _cell(c.title),
        _cell(c.preconditions),
        _cell(c.steps),
        _cell(c.expected),
        "-",
        "-",
        "-",
        "-",
        "-",
        "-",
        _cell(c.remark or (f"P{c.priority[-1]};{c.category}")),
    ]


def export_cases(cases: list[TestCase], fmt: str, base_name: str | None = None) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = base_name or f"功能测试用例_{stamp}"

    if fmt == "xlsx":
        path = OUTPUT_DIR / f"{name}.xlsx"
        _to_xlsx(cases, path)
    elif fmt == "md":
        path = OUTPUT_DIR / f"{name}.md"
        _to_md(cases, path)
    else:
        path = OUTPUT_DIR / f"{name}.csv"
        _to_csv(cases, path)
    return path


def _to_xlsx(cases: list[TestCase], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "功能测试用例"

    header_fill = PatternFill("solid", fgColor="D9D9D9")
    thin = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    header_font = Font(name="宋体", size=11, bold=True, color="000000")
    body_font = Font(name="宋体", size=11, color="000000")

    ws.append(V3_HEADERS)
    for col in range(1, len(V3_HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin

    for c in cases:
        ws.append(_row(c))

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(V3_HEADERS)):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin

    widths = [12, 14, 12, 12, 14, 28, 16, 18, 28, 32, 40, 32, 10, 10, 12, 12, 10, 10, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def _to_md(cases: list[TestCase], path: Path) -> None:
    esc = lambda s: _cell(s).replace("|", "\\|")
    lines = [
        "# 功能测试用例",
        "",
        f"> 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  ",
        f"> 用例数量：{len(cases)}  ",
        "> 格式：功能测试提示词 v3.0",
        "",
        "| " + " | ".join(V3_HEADERS) + " |",
        "|" + "|".join(["---"] * len(V3_HEADERS)) + "|",
    ]
    for c in cases:
        lines.append("| " + " | ".join(esc(x) for x in _row(c)) + " |")
    path.write_text("\n".join(lines), encoding="utf-8")


def _to_csv(cases: list[TestCase], path: Path) -> None:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(V3_HEADERS)
    for c in cases:
        w.writerow(_row(c))
    path.write_text(buf.getvalue(), encoding="utf-8-sig")
